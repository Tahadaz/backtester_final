from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - operator-facing dependency guard
    raise SystemExit("Missing dependency. Install python-pptx and Pillow before building the deck.") from exc

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SCREENS_DIR = ROOT / "docs/fundamentals-layer/assets/screens"
DEFAULT_OUTPUT = ROOT / "out/Module_Analyse_Fondamentale.pptx"

BLUE = RGBColor(31, 78, 121)
MID_BLUE = RGBColor(46, 117, 182)
LIGHT_BLUE = RGBColor(217, 225, 242)
DARK = RGBColor(28, 35, 45)
MUTED = RGBColor(86, 96, 112)
WHITE = RGBColor(255, 255, 255)
PAPER = RGBColor(246, 248, 251)

SCREENSHOTS = {
    "tearsheet": "01-tearsheet.png",
    "valuation": "02-valuation-models.png",
    "wacc": "03-wacc-buildup.png",
    "scenarios": "04-scenarios.png",
    "sensitivity": "05-sensitivity.png",
    "scoring": "06-scoring.png",
    "diagnostics": "07-diagnostics.png",
}


def _set_bg(slide, color: RGBColor = PAPER) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _text_box(slide, text: str, x: float, y: float, w: float, h: float, *, size: int = 18, color: RGBColor = DARK, bold: bool = False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def _title(slide, title: str, kicker: str | None = None) -> None:
    _text_box(slide, title, 0.55, 0.28, 10.0, 0.45, size=22, color=BLUE, bold=True)
    if kicker:
        _text_box(slide, kicker, 0.57, 0.76, 8.5, 0.32, size=9, color=MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.98), Inches(12.25), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.color.rgb = LIGHT_BLUE


def _footer(slide, number: int) -> None:
    _text_box(slide, "Module Analyse Fondamentale", 0.55, 7.14, 4.0, 0.22, size=8, color=MUTED)
    box = _text_box(slide, f"{number:02d}", 12.2, 7.12, 0.8, 0.24, size=8, color=MUTED)
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _fit_image(slide, image_path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(image_path) as img:
        px_w, px_h = img.size
    box_w = Inches(w)
    box_h = Inches(h)
    scale = min(int(box_w) / px_w, int(box_h) / px_h)
    pic_w = int(px_w * scale)
    pic_h = int(px_h * scale)
    left = int(Inches(x)) + int((int(box_w) - pic_w) / 2)
    top = int(Inches(y)) + int((int(box_h) - pic_h) / 2)
    slide.shapes.add_picture(str(image_path), left, top, width=pic_w, height=pic_h)


def _bullet_panel(slide, title: str, bullets: list[str], x: float = 10.0, y: float = 1.24, w: float = 2.72, h: float = 5.7) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = LIGHT_BLUE
    _text_box(slide, title, x + 0.18, y + 0.18, w - 0.36, 0.42, size=13, color=BLUE, bold=True)
    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.78), Inches(w - 0.34), Inches(h - 1.0))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(8)


def _screenshot_slide(prs: Presentation, number: int, title: str, image_path: Path, panel_title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _title(slide, title, "Vue issue du dashboard Signals / Fundamental")
    _fit_image(slide, image_path, 0.55, 1.22, 9.25, 5.85)
    _bullet_panel(slide, panel_title, bullets)
    _footer(slide, number)


def _title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, WHITE)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.18))
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.color.rgb = BLUE
    _text_box(slide, "Module Analyse Fondamentale", 0.72, 1.95, 11.2, 0.7, size=34, color=BLUE, bold=True)
    _text_box(slide, "Deck produit - scenarios, valorisation, scoring et garde-fous", 0.76, 2.75, 10.5, 0.4, size=17, color=DARK)
    _text_box(slide, "Cas d'usage: produire une note action exploitable, ancree sur le scenario de base, avec vues bear/base/bull pour le risque.", 0.78, 3.48, 10.4, 0.65, size=14, color=MUTED)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.78), Inches(4.65), Inches(4.3), Inches(0.07))
    accent.fill.solid()
    accent.fill.fore_color.rgb = MID_BLUE
    accent.line.color.rgb = MID_BLUE
    _text_box(slide, "Version demo - donnees live StockAnalysis ou fixture synthetique", 0.78, 6.83, 8.4, 0.3, size=9, color=MUTED)
    _footer(slide, 1)


def _text_slide(prs: Presentation, number: int, title: str, sections: list[tuple[str, list[str]]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide)
    _title(slide, title)
    x_positions = [0.65, 4.62, 8.59]
    for index, (section_title, bullets) in enumerate(sections):
        x = x_positions[index % len(x_positions)]
        y = 1.38 + (index // len(x_positions)) * 2.85
        _bullet_panel(slide, section_title, bullets, x=x, y=y, w=3.45, h=2.45)
    _footer(slide, number)


def _require_screens(screens_dir: Path) -> dict[str, Path]:
    paths = {key: screens_dir / file_name for key, file_name in SCREENSHOTS.items()}
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise SystemExit("Missing required screenshots:\n" + "\n".join(missing))
    return paths


def build_deck(screens_dir: Path, output: Path) -> Path:
    paths = _require_screens(screens_dir)
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _title_slide(prs)
    _screenshot_slide(
        prs,
        2,
        "1. Ticket de recherche",
        paths["tearsheet"],
        "Decision",
        ["Reco et objectif ancrés au scenario base.", "Upside, conviction et revision restent lisibles.", "Le prix courant conserve sa source et sa date."],
    )
    _screenshot_slide(
        prs,
        3,
        "2. Modeles de valorisation",
        paths["valuation"],
        "Lecture",
        ["Modele par modele avec confiance et poids.", "Reverse DCF reste diagnostic.", "Les avertissements expliquent les exclusions."],
    )
    _screenshot_slide(
        prs,
        4,
        "3. Build-up WACC",
        paths["wacc"],
        "Discipline",
        ["Taux sans risque, beta, ERP et cout de dette visibles.", "Le floor de cout des fonds propres est explicite.", "Les sources live et registry sont separees."],
    )
    _screenshot_slide(
        prs,
        5,
        "4. Scenarios bear/base/bull",
        paths["scenarios"],
        "Gouvernance",
        ["Probabilites stockees comme hypotheses.", "La valeur ponderee est separee de la reco.", "Le scenario consulte ne deplace pas le headline."],
    )
    _screenshot_slide(
        prs,
        6,
        "5. Sensibilites",
        paths["sensitivity"],
        "Risque",
        ["Matrices WACC / croissance terminale.", "Table exploitable pour comite d'investissement.", "Les zones extremes cadrent le stress-test."],
    )
    _screenshot_slide(
        prs,
        7,
        "6. Scoring fondamental",
        paths["scoring"],
        "Priorisation",
        ["Qualite, value, croissance, risque et cash-flow.", "Historique de piliers pour lire la trajectoire.", "Coverage et diagnostics expliquent les trous de donnees."],
    )
    _screenshot_slide(
        prs,
        8,
        "7. Diagnostics",
        paths["diagnostics"],
        "Audit",
        ["DuPont, accruals et signaux comptables.", "Confiance degradee quand les donnees manquent.", "Traçabilite jusqu'aux imports et hypotheses."],
    )
    _text_slide(
        prs,
        9,
        "Garde-fous integres",
        [
            ("Scenario", ["Headline force sur base.", "Scenario implicite marche = diagnostic.", "Base absent => NR."]),
            ("Hypotheses", ["Hierarchie desk/secteur/symbole.", "Probabilites normalisees a 100%.", "Warning si renormalisation."]),
            ("Qualite", ["Integrity checks bloquants.", "Modeles proxys plafonnes.", "Warnings visibles dans l'UI."]),
        ],
    )
    _text_slide(
        prs,
        10,
        "Feuille de route",
        [
            ("Court terme", ["Capture automatisee apres refresh.", "Templates de note action.", "Export PDF depuis la tear sheet."]),
            ("Moyen terme", ["Overrides analyste versionnes.", "Comparables sectoriels enrichis.", "Beta et courbes de taux historises."]),
            ("Comite", ["Pack screens + PPTX reproductible.", "Scenario policy auditable.", "Journal des changements de reco."]),
        ],
    )

    prs.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the French fundamentals product deck from captured dashboard screenshots.")
    parser.add_argument("--screens-dir", type=Path, default=DEFAULT_SCREENS_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    path = build_deck(args.screens_dir, args.output)
    print(path)


if __name__ == "__main__":
    main()
