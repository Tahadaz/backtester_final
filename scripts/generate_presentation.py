from pptx import Presentation
from pptx.util import Inches, Pt
import os, json

# Short, concise PPTX generator — outputs a 6-slide summary using extracted results
DOWNLOADS = os.path.expanduser(r"c:\Users\taha\Downloads")
IMG1 = os.path.join(DOWNLOADS, "backtestMC.png")
from pptx import Presentation
from pptx.util import Inches, Pt
import os, json

# Short, concise PPTX generator — outputs a 6-slide summary using extracted results
DOWNLOADS = os.path.expanduser(r"c:\Users\taha\Downloads")
IMG1 = os.path.join(DOWNLOADS, "backtestMC.png")
IMG2 = os.path.join(DOWNLOADS, "MC.png")
OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs", "presentations"))
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "weekly_report_2026-04-24.pptx")

prs = Presentation()

def add_bullet_slide(title_text, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title_text
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for b in bullets:
        p = body.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)
    return slide

# Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Weekly Report — Week ending 2026-04-24"
slide.placeholders[1].text = "Highlights: robustness layer, macro/factor study, app integration"

# Executive summary
add_bullet_slide("Executive Summary", [
    "Added per-signal backtest + Monte Carlo robustness panel to signal pages",
    "Completed macro/factor descriptive study for MASI (results extracted)",
    "Drafted app-level factor-layer design and analytics API/frontend contract",
])

# Key concise findings
SUMMARY_FILE = os.path.abspath(os.path.join(OUTPUT_DIR, "factor_summary.json"))
findings = []
if os.path.exists(SUMMARY_FILE):
    with open(SUMMARY_FILE, 'r', encoding='utf-8') as f:
        summary = json.load(f)
    if 'ATW' in summary:
        atw = {r['factor']: r for r in summary['ATW']}
        if 'us10y' in atw:
            u = atw['us10y']
            findings.append(f"ATW — US10Y: Spearman {u['spearman_corr']:.3f}, p={u['spearman_p']:.3f}")
        if 'vix' in atw:
            v = atw['vix']
            findings.append(f"ATW — VIX: Spearman {v['spearman_corr']:.3f}, p={v['spearman_p']:.3f}")
    if 'MASI20' in summary:
        for r in summary['MASI20']:
            if r['factor'] == 'brent':
                findings.append(f"MASI20 — Brent: Spearman {r['spearman_corr']:.3f}, p={r['spearman_p']:.3f}")
                break

if not findings:
    findings = ["Detailed results are available in docs/presentations/factor_summary.json"]

add_bullet_slide("Key Findings (concise)", findings)

# Work completed
add_bullet_slide("Work Completed This Week", [
    "Per-signal backtest + Monte Carlo robustness UI integrated on signal pages",
    "Macro/factor study executed; outputs saved under iLoveZIP_Create/factor_lab_outputs",
    "Docs and app-level factor-layer design drafted (docs/factor-layer)",
])

# Next steps
add_bullet_slide("Next Steps", [
    "Add any additional symbol numbers to the PPTX if needed",
    "Run unit tests for stats/alignment and ingest key macro series for validation",
    "Prepare PR with code + docs for review",
])

# Artifacts
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Artifacts"
body = slide.shapes.placeholders[1].text_frame
body.clear()
body.add_paragraph().text = f"Local analysis: c:\\Users\\taha\\Downloads\\iLoveZIP_Create"
body.add_paragraph().text = f"Extracted summary: docs/presentations/factor_summary.json"
body.add_paragraph().text = f"PPTX output: {OUTPUT_FILE}"

# Optional images on title slide if available
img_left = Inches(0.5)
img_top = Inches(4.0)
img_w = Inches(3.5)
if os.path.exists(IMG1):
    try:
        prs.slides[0].shapes.add_picture(IMG1, img_left, img_top, width=img_w)
    except Exception:
        pass
if os.path.exists(IMG2):
    try:
        prs.slides[0].shapes.add_picture(IMG2, img_left + img_w + Inches(0.2), img_top, width=img_w)
    except Exception:
        pass

prs.save(OUTPUT_FILE)
print("Saved:", OUTPUT_FILE)
