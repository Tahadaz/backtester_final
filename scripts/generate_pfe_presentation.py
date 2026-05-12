from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import requests
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FRONTEND_DIR = ROOT / "frontend"
OUT_DIR = ROOT / "docs" / "presentations"
ASSET_DIR = OUT_DIR / "pfe_assets"
SCREEN_DIR = ASSET_DIR / "screens"
CROP_DIR = ASSET_DIR / "crops"
OUTPUT_PPTX = OUT_DIR / "pfe_signal_dashboard_engineer_grade.pptx"
OUTPUT_NOTES = OUT_DIR / "pfe_signal_dashboard_speaker_notes.md"

FRONTEND_BASE = "http://localhost:3001"
API_BASE = "http://localhost:8000"
DEFAULT_SCREEN_SOURCE = Path(
    r"C:\Users\taha\AppData\Local\Temp\pfe_screens_20260511_001356"
)


WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)


NAVY = RGBColor(18, 42, 78)
NAVY_2 = RGBColor(31, 64, 112)
BLUE = RGBColor(30, 86, 180)
BLUE_LIGHT = RGBColor(232, 240, 255)
BG = RGBColor(248, 250, 252)
CARD = RGBColor(255, 255, 255)
LINE = RGBColor(218, 225, 234)
TEXT = RGBColor(20, 26, 38)
MUTED = RGBColor(96, 110, 130)
GREEN = RGBColor(0, 132, 88)
GREEN_BG = RGBColor(228, 248, 238)
RED = RGBColor(206, 41, 50)
RED_BG = RGBColor(255, 235, 238)
AMBER = RGBColor(176, 111, 0)
AMBER_BG = RGBColor(255, 246, 224)
GRAY_BG = RGBColor(241, 245, 249)


@dataclass
class LiveStats:
    data_actions: int = 79
    data_bars: int = 95649
    data_with_data_pct: str = "100%"
    latest_bar: str = "2026-05-08"
    dashboard_setups: int = 14
    dashboard_longs: int = 7
    dashboard_shorts: int = 7
    dashboard_median_er: str = "2.72%"
    lhm_price: str = "1 849,00"
    lhm_signal: str = "Vente"
    lhm_direction: str = "Short"
    lhm_action_er: str = "1.94%"
    lhm_hit_rate: str = "83.33%"
    lhm_n: int = 42
    lhm_holding: str = "21 j"
    evidence_trades: int = 358
    evidence_periods: int = 10
    lhm_edge: str = "Prouvé"


def pct(value: float | None) -> str:
    if value is None:
        return "--"
    return f"{value * 100:.2f}%"


def format_int(value: int | float | None) -> str:
    if value is None:
        return "--"
    return f"{int(round(value)):,}".replace(",", " ")


def fetch_live_stats() -> LiveStats:
    stats = LiveStats()
    try:
        catalog = requests.get(f"{API_BASE}/market-data/catalog", timeout=15).json()
        actions = [r for r in catalog if (r.get("asset_type") or "equity") == "equity"]
        with_data = [r for r in actions if r.get("has_canonical_data")]
        stats.data_actions = len(actions)
        stats.data_bars = sum(int(r.get("row_count") or 0) for r in actions)
        stats.data_with_data_pct = (
            f"{round(100 * len(with_data) / len(actions))}%"
            if actions
            else "--"
        )
        latest = max((r.get("data_as_of") or "" for r in actions), default="")
        if latest:
            stats.latest_bar = latest
    except Exception:
        pass

    try:
        dash = requests.get(f"{API_BASE}/dashboard/data/monthly", timeout=20).json()
        stocks = dash.get("stocks", [])
        best = [s for s in stocks if s.get("best_signal")]
        actionable = [
            s
            for s in best
            if (s["best_signal"].get("direction") in ("long", "short"))
            and s["best_signal"].get("bucket") not in (None, "hold")
        ]
        stats.dashboard_setups = len(actionable)
        stats.dashboard_longs = sum(
            1 for s in actionable if s["best_signal"].get("direction") == "long"
        )
        stats.dashboard_shorts = sum(
            1 for s in actionable if s["best_signal"].get("direction") == "short"
        )
        ers = sorted(
            [
                float(s["best_signal"]["action_expected_return_net"])
                for s in actionable
                if s["best_signal"].get("action_expected_return_net") is not None
            ]
        )
        if ers:
            mid = len(ers) // 2
            med = ers[mid] if len(ers) % 2 else (ers[mid - 1] + ers[mid]) / 2
            stats.dashboard_median_er = pct(med)
        lhm = next((s for s in stocks if s.get("symbol") == "LHM"), None)
        if lhm:
            bs = lhm.get("best_signal") or {}
            if lhm.get("last_price") is not None:
                stats.lhm_price = format_int(lhm.get("last_price"))
            stats.lhm_signal = str(bs.get("signal_label") or stats.lhm_signal)
            stats.lhm_direction = (
                "Short" if bs.get("direction") == "short" else "Long"
            )
            stats.lhm_action_er = pct(bs.get("action_expected_return_net"))
            stats.lhm_hit_rate = pct(bs.get("hit_rate"))
            stats.lhm_n = int(bs.get("n") or stats.lhm_n)
            if bs.get("fwd_horizon_bars") is not None:
                stats.lhm_holding = f"{int(bs['fwd_horizon_bars'])} j"
            stats.lhm_edge = "Prouvé" if bs.get("proven_edge_net") else "Watch"
    except Exception:
        pass
    return stats


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SCREEN_DIR.mkdir(parents=True, exist_ok=True)
    CROP_DIR.mkdir(parents=True, exist_ok=True)


SCREENSHOT_URLS = {
    "dashboard_v1.png": f"{FRONTEND_BASE}/v1",
    "data_page.png": f"{FRONTEND_BASE}/data",
    "signals_lhm_technique.png": (
        f"{FRONTEND_BASE}/signals?symbol=LHM&horizon=monthly&tab=technique"
    ),
    "signals_lhm_evidence.png": (
        f"{FRONTEND_BASE}/signals?symbol=LHM&horizon=monthly&tab=evidence"
    ),
    "signals_lhm_wfo.png": (
        f"{FRONTEND_BASE}/signals?symbol=LHM&horizon=monthly&tab=wfo"
    ),
    "analytics_page.png": f"{FRONTEND_BASE}/analytics",
}


def run_playwright_screenshot(name: str, url: str) -> bool:
    target = SCREEN_DIR / name
    try:
        if os.name == "nt":
            cmd = (
                'npx playwright screenshot --full-page --wait-for-timeout=6500 '
                f'--viewport-size=1600,1000 "{url}" "{target}"'
            )
            subprocess.run(cmd, cwd=FRONTEND_DIR, check=True, timeout=120, shell=True)
        else:
            cmd = [
                "npx",
                "playwright",
                "screenshot",
                "--full-page",
                "--wait-for-timeout=6500",
                "--viewport-size=1600,1000",
                url,
                str(target),
            ]
            subprocess.run(cmd, cwd=FRONTEND_DIR, check=True, timeout=120)
        return target.exists()
    except Exception:
        return False


def collect_screenshots(refresh: bool = False) -> None:
    source = Path(os.environ.get("PFE_SCREEN_SOURCE", str(DEFAULT_SCREEN_SOURCE)))
    for name, url in SCREENSHOT_URLS.items():
        target = SCREEN_DIR / name
        if target.exists() and not refresh:
            continue
        source_file = source / name
        if source_file.exists() and not refresh:
            shutil.copy2(source_file, target)
            continue
        ok = run_playwright_screenshot(name, url)
        if not ok and source_file.exists():
            shutil.copy2(source_file, target)


def crop_image(name: str, box: tuple[int, int, int, int]) -> Path:
    src = SCREEN_DIR / name
    out = CROP_DIR / f"{Path(name).stem}_crop.png"
    img = Image.open(src).convert("RGB")
    w, h = img.size
    left, top, right, bottom = box
    left = max(0, min(left, w - 1))
    top = max(0, min(top, h - 1))
    right = max(left + 1, min(right, w))
    bottom = max(top + 1, min(bottom, h))
    img.crop((left, top, right, bottom)).save(out, quality=95)
    return out


def make_crops() -> dict[str, Path]:
    crops = {
        "dashboard_top": crop_image("dashboard_v1.png", (145, 70, 1465, 1020)),
        "dashboard_table": crop_image("dashboard_v1.png", (170, 940, 1460, 1910)),
        "dashboard_lhm_row": crop_image("dashboard_v1.png", (170, 1445, 1460, 1745)),
        "data_overview": crop_image("data_page.png", (150, 70, 1465, 1145)),
        "signals_technique": crop_image("signals_lhm_technique.png", (230, 60, 1600, 955)),
        "signals_evidence": crop_image("signals_lhm_evidence.png", (240, 65, 1600, 980)),
        "signals_wfo": crop_image("signals_lhm_wfo.png", (240, 65, 1600, 980)),
        "analytics": crop_image("analytics_page.png", (225, 55, 1590, 940)),
    }
    return crops


def set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def set_line(shape, color: RGBColor, width: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: RGBColor = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "Aptos",
    italic: bool = False,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_rich_lines(
    slide,
    lines: Iterable[tuple[str, int, RGBColor, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    spacing: int = 4,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(spacing)
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return box


def add_card(slide, x: float, y: float, w: float, h: float, fill: RGBColor = CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, LINE, 0.9)
    return shape


def add_pill(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float = 0.32,
    *,
    fill: RGBColor = BLUE_LIGHT,
    color: RGBColor = BLUE,
    size: int = 10,
    bold: bool = True,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, fill, 0.1)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def add_footer(slide, number: int) -> None:
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.2), WIDE_W, Inches(0.01)
    )
    set_fill(line, LINE)
    add_text(
        slide,
        "PFE 2026 - Backtest Signal Engine",
        0.55,
        7.24,
        4.0,
        0.18,
        size=8,
        color=MUTED,
    )
    add_text(
        slide,
        f"{number:02d}",
        12.45,
        7.20,
        0.4,
        0.2,
        size=8,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_header(slide, title: str, subtitle: str, number: int) -> None:
    add_text(slide, title, 0.55, 0.34, 8.7, 0.36, size=22, color=TEXT, bold=True)
    add_text(slide, subtitle, 0.56, 0.78, 10.9, 0.32, size=10, color=MUTED)
    add_footer(slide, number)


def add_image_fit(slide, image_path: Path, x: float, y: float, w: float, h: float):
    return slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_callout(slide, text: str, x: float, y: float, w: float = 1.5, *, tone: str = "blue"):
    if tone == "green":
        fill, color = GREEN_BG, GREEN
    elif tone == "red":
        fill, color = RED_BG, RED
    elif tone == "amber":
        fill, color = AMBER_BG, AMBER
    else:
        fill, color = BLUE_LIGHT, BLUE
    return add_pill(slide, text, x, y, w, 0.34, fill=fill, color=color, size=9, bold=True)


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = BLUE):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(1.7)
    conn.line.end_arrowhead = True
    return conn


def add_section_label(slide, text: str, x: float, y: float, w: float):
    add_text(
        slide,
        text.upper(),
        x,
        y,
        w,
        0.22,
        size=8,
        color=MUTED,
        bold=True,
    )


def add_stat_card(slide, label: str, value: str, sub: str, x: float, y: float, w: float):
    add_card(slide, x, y, w, 0.9, CARD)
    add_text(slide, label.upper(), x + 0.12, y + 0.12, w - 0.24, 0.15, size=7, color=MUTED, bold=True)
    add_text(slide, value, x + 0.12, y + 0.34, w - 0.24, 0.25, size=17, color=TEXT, bold=True)
    add_text(slide, sub, x + 0.12, y + 0.64, w - 0.24, 0.16, size=7.5, color=MUTED)


def blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDE_W, WIDE_H)
    set_fill(bg, BG)
    bg.line.color.rgb = BG
    return slide


def slide_title(prs: Presentation, crops: dict[str, Path]) -> None:
    slide = blank(prs)
    left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.25), WIDE_H)
    set_fill(left, NAVY)
    add_text(slide, "PFE - Projet de fin d'études", 0.55, 0.55, 3.5, 0.25, size=11, color=RGBColor(205, 219, 245), bold=True)
    add_text(slide, "Solution d'aide à la décision\nen trading actions", 0.55, 1.35, 4.25, 1.35, size=28, color=RGBColor(255, 255, 255), bold=True)
    add_text(slide, "Dashboard quotidien - Signal expliqué - Validation hors échantillon", 0.58, 3.08, 3.95, 0.45, size=12, color=RGBColor(220, 230, 246))
    add_pill(slide, "Dashboard", 0.58, 4.15, 1.3, fill=RGBColor(38, 76, 132), color=RGBColor(255, 255, 255))
    add_pill(slide, "Signal Engine", 2.02, 4.15, 1.45, fill=RGBColor(38, 76, 132), color=RGBColor(255, 255, 255))
    add_pill(slide, "Finance quantitative", 0.58, 4.60, 1.95, fill=RGBColor(38, 76, 132), color=RGBColor(255, 255, 255))
    add_text(slide, "Réalisé par : Taha Dazine", 0.58, 6.65, 3.4, 0.25, size=10, color=RGBColor(220, 230, 246))
    add_card(slide, 5.6, 0.62, 7.15, 5.85, RGBColor(255, 255, 255))
    add_image_fit(slide, crops["dashboard_top"], 5.75, 0.78, 6.85, 5.45)
    add_text(slide, "Version actuelle de l'application - captures live locales", 5.78, 6.60, 5.8, 0.2, size=8, color=MUTED)
    add_footer(slide, 1)


def slide_context(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Contexte métier du desk", "La valeur n'est pas d'avoir plus d'indicateurs, mais de savoir où regarder en premier.", 2)
    cols = [
        ("Constat", "Beaucoup d'actions à suivre\nplusieurs horizons\nplusieurs signaux par titre", BLUE_LIGHT, BLUE),
        ("Risque", "Lecture dispersée du marché\nhiérarchisation difficile\nopportunités faciles à manquer", RED_BG, RED),
        ("Besoin", "Une shortlist quotidienne\nun signal centralisé\nune justification immédiate", GREEN_BG, GREEN),
    ]
    x = 0.75
    for i, (title, body, fill, color) in enumerate(cols):
        add_card(slide, x + i * 4.15, 1.65, 3.55, 3.55, CARD)
        add_pill(slide, str(i + 1), x + i * 4.15 + 0.25, 1.92, 0.42, fill=fill, color=color, size=12)
        add_text(slide, title, x + i * 4.15 + 0.82, 1.90, 2.35, 0.28, size=17, color=TEXT, bold=True)
        add_text(slide, body, x + i * 4.15 + 0.35, 2.55, 2.8, 1.3, size=13, color=TEXT)
    add_callout(slide, "Question centrale : quels titres méritent une analyse prioritaire aujourd'hui ?", 2.15, 5.80, 8.95)


def slide_objective(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Objectif et périmètre du projet", "Construire une chaîne explicable entre données de marché, signal et lecture décisionnelle.", 3)
    steps = [
        ("Données", "OHLCV validées\ncatalogue, fraîcheur"),
        ("Signaux", "indicateurs\nOOS + robustesse"),
        ("Dashboard", "tri par horizon\nshortlist quotidienne"),
        ("Preuve", "détails signal\ntrades OOS, WFO"),
    ]
    x0 = 0.85
    for i, (t, b) in enumerate(steps):
        x = x0 + i * 3.1
        add_card(slide, x, 2.05, 2.35, 2.2, CARD)
        add_pill(slide, t, x + 0.25, 2.35, 1.85, fill=BLUE_LIGHT, color=BLUE, size=11)
        add_text(slide, b, x + 0.28, 2.95, 1.8, 0.75, size=11, color=TEXT, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.4, 3.15, x + 2.92, 3.15)
    add_card(slide, 1.1, 5.05, 11.1, 0.95, RGBColor(255, 255, 255))
    add_text(slide, "Positionnement", 1.35, 5.22, 1.6, 0.22, size=10, color=BLUE, bold=True)
    add_text(
        slide,
        "L'application ne remplace pas le jugement du trader : elle organise l'information, documente la preuve et réduit le temps nécessaire pour identifier les dossiers à approfondir.",
        2.65,
        5.20,
        8.75,
        0.42,
        size=11,
        color=TEXT,
    )


def slide_data(prs: Presentation, crops: dict[str, Path], stats: LiveStats) -> None:
    slide = blank(prs)
    add_header(slide, "Fondation données", "Les signaux ne sont fiables que si la base OHLCV est contrôlée, fraîche et reproductible.", 4)
    add_image_fit(slide, crops["data_overview"], 0.55, 1.25, 8.15, 5.25)
    add_stat_card(slide, "Actions", str(stats.data_actions), "catalogue actions", 9.15, 1.35, 3.25)
    add_stat_card(slide, "Daily bars", format_int(stats.data_bars), "historique exploitable", 9.15, 2.42, 3.25)
    add_stat_card(slide, "Avec données", stats.data_with_data_pct, "instruments couverts", 9.15, 3.49, 3.25)
    add_stat_card(slide, "Dernier bar", stats.latest_bar, "base live de référence", 9.15, 4.56, 3.25)
    add_callout(slide, "Imports Excel + sources marché + contrôles de fraîcheur", 9.2, 5.8, 3.1)


def slide_finance(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Raisonnement financier", "Chaque famille d'indicateurs traduit une hypothèse économique différente.", 5)
    cards = [
        ("Tendance", "SMA / EMA / PSAR", "Le prix est-il durablement orienté ?", "Direction dominante", BLUE_LIGHT, BLUE),
        ("Momentum", "MACD / ROC / ADX", "Le mouvement accélère-t-il ?", "Force du mouvement", AMBER_BG, AMBER),
        ("Oscillation", "RSI / Stochastic / CCI", "Le marché est-il en excès ?", "Surachat / survente", RED_BG, RED),
        ("Volume", "OBV / CMF / VWAP", "Les flux confirment-ils le prix ?", "Accumulation / distribution", GREEN_BG, GREEN),
    ]
    positions = [(0.7, 1.55), (6.95, 1.55), (0.7, 4.05), (6.95, 4.05)]
    for (title, inds, question, role, fill, color), (x, y) in zip(cards, positions):
        add_card(slide, x, y, 5.65, 1.85, CARD)
        add_pill(slide, title, x + 0.25, y + 0.20, 1.45, fill=fill, color=color)
        add_text(slide, inds, x + 1.9, y + 0.22, 3.1, 0.25, size=10, color=MUTED, bold=True)
        add_text(slide, question, x + 0.28, y + 0.72, 4.8, 0.28, size=13, color=TEXT, bold=True)
        add_text(slide, role, x + 0.28, y + 1.15, 4.7, 0.28, size=10, color=MUTED)
    add_callout(slide, "Le signal global est un consensus : il combine plusieurs lectures, pas plusieurs copies de la même idée.", 1.2, 6.35, 10.9)


def slide_pipeline(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Pipeline du signal", "Le moteur évite de choisir directement le meilleur indicateur observé sur l'historique.", 6)
    stages = [
        ("A", "Candidats", "grilles cohérentes par horizon"),
        ("B", "OOS", "évaluation walk-forward"),
        ("C", "Robustesse", "Sharpe, stabilité, drawdown"),
        ("D", "Survivants", "filtrage des faibles variantes"),
        ("E", "Redondance", "corrélation max 0.85"),
        ("F", "Signal courant", "dernier bar disponible"),
        ("G", "Ensemble", "consensus pondéré"),
    ]
    x = 0.45
    for i, (letter, title, body) in enumerate(stages):
        w = 1.65
        add_card(slide, x + i * 1.83, 2.05, w, 2.55, CARD)
        add_pill(slide, letter, x + i * 1.83 + 0.52, 2.28, 0.5, fill=BLUE_LIGHT, color=BLUE, size=13)
        add_text(slide, title, x + i * 1.83 + 0.12, 2.90, w - 0.24, 0.27, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + i * 1.83 + 0.13, 3.35, w - 0.26, 0.65, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            add_arrow(slide, x + i * 1.83 + w, 3.28, x + (i + 1) * 1.83 - 0.08, 3.28, MUTED)
    add_card(slide, 1.15, 5.28, 11.0, 0.85, RGBColor(255, 255, 255))
    add_text(slide, "Principe d'ingénierie", 1.4, 5.46, 2.0, 0.2, size=10, color=BLUE, bold=True)
    add_text(slide, "Chaque score affiché dans l'interface reste traçable : famille -> représentants -> fenêtres OOS -> trades observés.", 3.35, 5.44, 7.9, 0.25, size=11, color=TEXT)


def slide_math(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Mathématiques essentielles", "Les formules restent simples afin que le score soit auditable par un utilisateur métier.", 7)
    formulas = [
        ("Rendement net", "R_net(t+1) = p_t r(t+1) - k |p_t - p_{t-1}|", "Le signal est pénalisé lorsqu'il tourne trop souvent."),
        ("Sharpe annualisé", "Sharpe = mean(R_net) / std(R_net) x sqrt(252)", "Mesure du rendement ajusté du risque."),
        ("Robustesse", "0.35 Sharpe + 0.30 Stabilité + 0.20 Consistance + 0.15 Drawdown", "Une variante ne gagne pas grâce à une seule bonne métrique."),
        ("Ensemble", "S = sum R(v) s(v) / sum R(v)", "Les représentants fiables pèsent davantage dans le consensus."),
    ]
    for i, (title, formula, body) in enumerate(formulas):
        x = 0.85 if i % 2 == 0 else 6.95
        y = 1.5 + (i // 2) * 2.35
        add_card(slide, x, y, 5.5, 1.75, CARD)
        add_text(slide, title, x + 0.25, y + 0.22, 2.4, 0.22, size=11, color=BLUE, bold=True)
        add_text(slide, formula, x + 0.25, y + 0.62, 4.85, 0.32, size=13, color=TEXT, bold=True)
        add_text(slide, body, x + 0.25, y + 1.10, 4.85, 0.28, size=9.5, color=MUTED)
    add_callout(slide, "Objectif : mesurer la qualité informative d'un signal, pas optimiser une promesse de performance.", 1.3, 6.45, 10.6)


def slide_oos(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Validation hors échantillon", "Le moteur juge les signaux sur des périodes postérieures à la fenêtre de contexte.", 8)
    y = 2.55
    add_text(slide, "Fenêtre 1", 0.85, 1.65, 1.1, 0.2, size=9, color=MUTED, bold=True)
    add_card(slide, 0.85, y, 3.2, 0.55, BLUE_LIGHT)
    add_text(slide, "Train / contexte", 1.65, y + 0.15, 1.6, 0.15, size=9, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 4.05, y, 1.6, 0.55, GREEN_BG)
    add_text(slide, "Test OOS", 4.38, y + 0.15, 0.9, 0.15, size=9, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 5.9, y + 0.27, 6.5, y + 0.27, MUTED)
    add_text(slide, "Fenêtre 2", 6.75, 1.65, 1.1, 0.2, size=9, color=MUTED, bold=True)
    add_card(slide, 6.75, y, 3.2, 0.55, BLUE_LIGHT)
    add_text(slide, "Train / contexte", 7.55, y + 0.15, 1.6, 0.15, size=9, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, 9.95, y, 1.6, 0.55, GREEN_BG)
    add_text(slide, "Test OOS", 10.28, y + 0.15, 0.9, 0.15, size=9, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    bullets = [
        ("Causalité", "les indicateurs utilisent uniquement le passé connu à la date t."),
        ("Coûts", "chaque changement de position paie un coût de transaction."),
        ("Stabilité", "un signal doit fonctionner sur plusieurs fenêtres, pas seulement une période chanceuse."),
        ("Horizon", "court, moyen et long ont des géométries de train/test différentes."),
    ]
    for i, (t, b) in enumerate(bullets):
        x = 0.9 + (i % 2) * 6.0
        yy = 4.05 + (i // 2) * 1.05
        add_card(slide, x, yy, 5.45, 0.72, CARD)
        add_text(slide, t, x + 0.2, yy + 0.13, 1.35, 0.18, size=10, color=BLUE, bold=True)
        add_text(slide, b, x + 1.45, yy + 0.12, 3.65, 0.25, size=9.5, color=TEXT)


def slide_dashboard(prs: Presentation, crops: dict[str, Path], stats: LiveStats) -> None:
    slide = blank(prs)
    add_header(slide, "Dashboard quotidien", "Le point d'entrée du desk : scanner vite, filtrer, puis ouvrir la preuve.", 9)
    add_image_fit(slide, crops["dashboard_top"], 0.55, 1.25, 8.95, 5.55)
    add_stat_card(slide, "Setups", str(stats.dashboard_setups), "après filtres", 9.85, 1.38, 2.75)
    add_stat_card(slide, "Long / Short", f"{stats.dashboard_longs} / {stats.dashboard_shorts}", "équilibre directionnel", 9.85, 2.45, 2.75)
    add_stat_card(slide, "Action E[R]", stats.dashboard_median_er, "médiane nette", 9.85, 3.52, 2.75)
    add_callout(slide, "Le dashboard trie par retour attendu et statut Edge", 9.9, 4.78, 2.65)
    add_callout(slide, "La ligne n'est pas une décision finale : elle ouvre l'analyse", 9.9, 5.35, 2.65, tone="amber")


def slide_lhm_case(prs: Presentation, crops: dict[str, Path], stats: LiveStats) -> None:
    slide = blank(prs)
    add_header(slide, "Cas concret : LHM", "La valeur du dashboard est de transformer une ligne en question analytique précise.", 10)
    add_image_fit(slide, crops["dashboard_lhm_row"], 0.75, 1.35, 11.85, 2.75)
    add_arrow(slide, 11.1, 3.25, 10.3, 2.35, GREEN)
    add_callout(slide, "Edge prouvé", 10.45, 3.35, 1.5, tone="green")
    metrics = [
        ("Signal", stats.lhm_signal, stats.lhm_direction, RED_BG, RED),
        ("Action E[R] net", stats.lhm_action_er, "après coûts", GREEN_BG, GREEN),
        ("Hit rate", stats.lhm_hit_rate, f"n={stats.lhm_n}", BLUE_LIGHT, BLUE),
        ("Horizon", stats.lhm_holding, "open-to-open", AMBER_BG, AMBER),
    ]
    for i, (label, value, sub, fill, color) in enumerate(metrics):
        add_card(slide, 0.75 + i * 3.0, 4.55, 2.55, 1.25, CARD)
        add_text(slide, label.upper(), 0.93 + i * 3.0, 4.72, 2.1, 0.18, size=7.5, color=MUTED, bold=True)
        add_text(slide, value, 0.93 + i * 3.0, 5.02, 2.1, 0.25, size=17, color=color, bold=True)
        add_text(slide, sub, 0.93 + i * 3.0, 5.36, 2.1, 0.18, size=8.5, color=MUTED)
    add_text(slide, "Lecture : un signal de vente avec Action E[R] positif signifie que, historiquement, le titre a baissé après des signaux comparables.", 1.15, 6.25, 10.8, 0.26, size=11, color=TEXT, bold=True)


def slide_technique(prs: Presentation, crops: dict[str, Path]) -> None:
    slide = blank(prs)
    add_header(slide, "Justification technique du signal", "La page Signaux explique pourquoi un titre ressort dans le dashboard.", 11)
    add_image_fit(slide, crops["signals_technique"], 0.55, 1.22, 12.2, 5.75)
    add_callout(slide, "Prix + marqueurs", 2.1, 2.05, 1.45)
    add_arrow(slide, 2.75, 2.42, 4.0, 3.85)
    add_callout(slide, "Familles techniques", 10.45, 1.65, 1.65)
    add_arrow(slide, 10.95, 2.0, 11.55, 2.55)
    add_callout(slide, "Scores SE / WFO", 8.85, 0.95, 1.55, tone="red")


def slide_evidence(prs: Presentation, crops: dict[str, Path], stats: LiveStats) -> None:
    slide = blank(prs)
    add_header(slide, "Preuve statistique OOS", "La justification relie le verdict courant à des trades observés hors échantillon.", 12)
    add_image_fit(slide, crops["signals_evidence"], 0.55, 1.20, 12.2, 5.75)
    add_callout(slide, f"{stats.evidence_trades} trades échantillonnés", 10.05, 1.55, 2.0, tone="blue")
    add_callout(slide, f"{stats.evidence_periods} périodes OOS", 10.15, 2.05, 1.75, tone="green")
    add_callout(slide, "Action nette > 0", 1.25, 2.85, 1.65, tone="green")
    add_arrow(slide, 2.05, 3.22, 3.05, 3.05, GREEN)


def slide_wfo(prs: Presentation, crops: dict[str, Path]) -> None:
    slide = blank(prs)
    add_header(slide, "Preuve Walk-Forward", "La WFO vérifie si les paramètres et catégories conservent une qualité hors échantillon.", 13)
    add_image_fit(slide, crops["signals_wfo"], 0.55, 1.18, 12.2, 5.78)
    add_callout(slide, "Consensus WFO", 1.15, 1.95, 1.55)
    add_callout(slide, "WFE + robustesse", 8.6, 2.05, 1.8, tone="amber")
    add_callout(slide, "Configuration train/test", 7.1, 5.48, 2.1, tone="blue")


def slide_edge_ticket(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Edge, ticket et blotter", "La couche décisionnelle transforme une preuve en préparation d'action, pas en exécution automatique.", 14)
    steps = [
        ("Signal courant", "bucket : achat / vente\nsource : SE ou WFO"),
        ("Edge proof", "n >= 30\nWilson > 50%\nMC + label shuffle"),
        ("Panier", "sélection des titres\nlong-only par défaut"),
        ("Sizing", "HRP + Kelly shrink\nplafonds titre/secteur"),
        ("Blotter", "BUY / HOLD / REDUCE\nEXIT / WATCH"),
    ]
    for i, (title, body) in enumerate(steps):
        x = 0.55 + i * 2.55
        add_card(slide, x, 2.15, 2.05, 2.2, CARD)
        add_pill(slide, title, x + 0.17, 2.40, 1.7, fill=BLUE_LIGHT, color=BLUE, size=9)
        add_text(slide, body, x + 0.22, 3.00, 1.55, 0.75, size=9, color=TEXT, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.08, 3.23, x + 2.45, 3.23, MUTED)
    add_card(slide, 1.05, 5.25, 11.1, 0.8, CARD)
    add_text(slide, "Garde-fou", 1.32, 5.44, 1.2, 0.2, size=10, color=RED, bold=True)
    add_text(slide, "Les niveaux d'entrée, stop et objectif sont des références pour la prochaine séance ; l'application ne route pas d'ordres et ne remplace pas la validation humaine.", 2.35, 5.42, 8.7, 0.25, size=10.5, color=TEXT)


def slide_architecture(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Architecture technique", "La plateforme sépare interface, contrats API, traitements longs et persistance.", 15)
    layers = [
        ("Frontend", "Next.js 16\nReact 19\nTypeScript\nSWR + Tailwind", 0.8, 1.65, BLUE_LIGHT, BLUE),
        ("API", "FastAPI\nPydantic\nSQLAlchemy\nAlembic", 3.25, 1.65, GREEN_BG, GREEN),
        ("Workers", "RQ + Redis\nsignal engine\nWFO\nsnapshots", 5.7, 1.65, AMBER_BG, AMBER),
        ("Storage", "PostgreSQL\nMinIO/S3\nparquet\nartefacts", 8.15, 1.65, GRAY_BG, NAVY_2),
        ("Quant core", "Python package\nindicateurs\nOOS / Edge\nportfolio", 10.6, 1.65, BLUE_LIGHT, BLUE),
    ]
    for i, (title, body, x, y, fill, color) in enumerate(layers):
        add_card(slide, x, y, 1.95, 3.05, CARD)
        add_pill(slide, title, x + 0.15, y + 0.22, 1.65, fill=fill, color=color, size=9)
        add_text(slide, body, x + 0.18, y + 0.85, 1.55, 1.25, size=10, color=TEXT, align=PP_ALIGN.CENTER)
        if i < len(layers) - 1:
            add_arrow(slide, x + 1.98, y + 1.55, x + 2.33, y + 1.55, MUTED)
    add_card(slide, 0.9, 5.45, 11.5, 0.65, CARD)
    add_text(slide, "Déploiement", 1.15, 5.62, 1.25, 0.18, size=9, color=BLUE, bold=True)
    add_text(slide, "Docker Compose orchestre frontend, API, workers, Postgres, Redis et MinIO ; Caddy sert de proxy en production.", 2.25, 5.60, 8.9, 0.22, size=10, color=TEXT)


def slide_limits(prs: Presentation) -> None:
    slide = blank(prs)
    add_header(slide, "Limites et perspectives", "La conclusion doit rester rigoureuse : l'outil aide à lire le marché, il ne le prédit pas parfaitement.", 16)
    items = [
        ("Limites actuelles", "Pas d'exécution automatique\npas de données intraday\nun edge historique peut se dégrader", RED_BG, RED),
        ("Perspectives", "Événements de marché\nfacteurs macro plus complets\nmonitoring de dérive des signaux", BLUE_LIGHT, BLUE),
        ("Contribution", "Pipeline reproductible\ndashboard décisionnel\npreuve explicable par titre", GREEN_BG, GREEN),
    ]
    for i, (title, body, fill, color) in enumerate(items):
        x = 0.9 + i * 4.1
        add_card(slide, x, 1.75, 3.4, 3.05, CARD)
        add_pill(slide, title, x + 0.35, 2.05, 2.2, fill=fill, color=color, size=10)
        add_text(slide, body, x + 0.38, 2.75, 2.55, 0.9, size=12, color=TEXT, align=PP_ALIGN.CENTER)
    add_text(slide, "Message final", 1.05, 5.55, 1.4, 0.2, size=10, color=BLUE, bold=True)
    add_text(slide, "La valeur du projet est de réduire le temps entre données brutes, détection d'une opportunité et justification méthodologique.", 2.25, 5.53, 9.55, 0.28, size=13, color=TEXT, bold=True)


def build_deck(crops: dict[str, Path], stats: LiveStats) -> Presentation:
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    slide_title(prs, crops)
    slide_context(prs)
    slide_objective(prs)
    slide_data(prs, crops, stats)
    slide_finance(prs)
    slide_pipeline(prs)
    slide_math(prs)
    slide_oos(prs)
    slide_dashboard(prs, crops, stats)
    slide_lhm_case(prs, crops, stats)
    slide_technique(prs, crops)
    slide_evidence(prs, crops, stats)
    slide_wfo(prs, crops)
    slide_edge_ticket(prs)
    slide_architecture(prs)
    slide_limits(prs)
    return prs


NOTES = [
    ("Solution d'aide à la décision", "0:00 - 0:45", "Introduire le projet comme une solution d'aide à la décision pour un desk actions. Le message important est que l'application ne se limite pas à afficher des indicateurs : elle relie données de marché, moteur de signaux, preuves statistiques et lecture opérationnelle dans un même parcours."),
    ("Contexte métier", "0:45 - 1:35", "Décrire le problème quotidien : un analyste doit suivre beaucoup de titres, plusieurs horizons et plusieurs familles de signaux. Sans priorisation, le risque est de passer du temps sur les mauvais dossiers ou de manquer une opportunité. Le dashboard répond à ce besoin de tri initial."),
    ("Objectif et périmètre", "1:35 - 2:25", "Présenter le périmètre de façon académique : collecte et contrôle des données, génération de signaux, dashboard de synthèse, puis justification détaillée. Préciser que le système ne prend pas la décision à la place de l'utilisateur ; il structure l'information et documente la preuve."),
    ("Fondation données", "2:25 - 3:20", "Montrer que la qualité des signaux dépend d'abord de la base OHLCV. Les chiffres live donnent la crédibilité de la démonstration : couverture actions, nombre de barres historiques et date du dernier bar disponible. Insister sur la fraîcheur et la reproductibilité des données."),
    ("Raisonnement financier", "3:20 - 4:20", "Expliquer les familles d'indicateurs comme des hypothèses financières complémentaires. La tendance répond à la direction, le momentum à la force, les oscillateurs aux excès de marché, et le volume à la confirmation des flux. Le consensus évite de dépendre d'une seule lecture fragile."),
    ("Pipeline du signal", "4:20 - 5:20", "Décrire le pipeline comme une chaîne de sélection robuste. Le moteur part de variantes candidates, les teste hors échantillon, filtre les variantes faibles, limite la redondance, puis agrège les représentants survivants. L'objectif est de réduire le sur-apprentissage historique."),
    ("Mathématiques essentielles", "5:20 - 6:25", "Expliquer les formules sans entrer dans une démonstration lourde. Le rendement net intègre les coûts de transaction, le Sharpe mesure le rendement ajusté du risque, la robustesse combine plusieurs critères, et l'ensemble pondère davantage les variantes fiables."),
    ("Validation hors échantillon", "6:25 - 7:20", "Insister sur la logique walk-forward : on observe une fenêtre de contexte, puis on teste sur une période postérieure. Cette discipline respecte la causalité et évite de juger un signal uniquement sur la période qui a servi à le choisir."),
    ("Dashboard quotidien", "7:20 - 8:20", "Présenter le dashboard comme le point d'entrée. L'utilisateur lit rapidement les setups actionnables, le sens long ou short, le retour attendu net, le hit rate et le statut Edge. Le but est de décider quels titres méritent une analyse approfondie."),
    ("Cas LHM", "8:20 - 9:15", "Utiliser LHM comme fil rouge. Expliquer qu'un signal Vente avec Action E[R] net positif signifie que, dans les observations comparables, une position short ou une baisse du titre a produit une espérance positive après coûts. La ligne ouvre l'analyse, elle ne conclut pas seule."),
    ("Justification technique", "9:15 - 10:10", "Montrer que la page Signaux replace la décision dans son contexte graphique. Les marqueurs buy, sell, short et cover expliquent le comportement historique du signal. Le panneau latéral montre les familles techniques et les sources Signal Engine ou WFO."),
    ("Preuve statistique OOS", "10:10 - 11:05", "Expliquer l'onglet Evidence : il documente les trades hors échantillon qui ressemblent au signal courant. Les métriques Action E[R], hit rate, holding et table de trades permettent de vérifier si le verdict actuel possède un précédent statistique exploitable."),
    ("Preuve Walk-Forward", "11:05 - 11:55", "Présenter la WFO comme un test de transférabilité. Les paramètres et catégories sélectionnés dans une fenêtre doivent conserver une qualité sur la fenêtre suivante. Les scores WFE et robustesse sont des diagnostics de stabilité, pas une garantie de performance future."),
    ("Edge, ticket et blotter", "11:55 - 12:45", "Décrire la couche décisionnelle : le système transforme un signal justifié en préparation d'action. Les garde-fous incluent taille d'échantillon, Wilson, Monte Carlo, label shuffle et coûts. Le blotter prépare BUY, HOLD, REDUCE, EXIT ou WATCH sans exécution automatique."),
    ("Architecture technique", "12:45 - 13:45", "Donner une lecture technique courte : Next.js et React pour l'interface, FastAPI pour les contrats backend, workers RQ pour les calculs longs, PostgreSQL pour la persistance, Redis pour la file, MinIO pour les artefacts et quant_core pour la logique financière."),
    ("Limites et perspectives", "13:45 - 15:00", "Conclure avec rigueur. Le projet accélère la lecture du marché et rend la justification traçable, mais un edge historique peut se dégrader et l'absence de données intraday limite certaines décisions. Les perspectives portent sur les événements de marché, les facteurs macro et le monitoring de dérive."),
]


def write_notes(stats: LiveStats) -> None:
    lines = [
        "# Notes orateur - Présentation PFE",
        "",
        f"Données live intégrées : {stats.data_actions} actions, {format_int(stats.data_bars)} barres, dernier bar {stats.latest_bar}.",
        f"Cas LHM : {stats.lhm_signal}, {stats.lhm_direction}, Action E[R] net {stats.lhm_action_er}, hit rate {stats.lhm_hit_rate}, n={stats.lhm_n}.",
        "",
    ]
    lines.extend([
        "Plan de parole conseillé : viser 14 à 15 minutes, en gardant 1 à 2 minutes pour les questions.",
        "",
    ])
    for i, (title, timing, note) in enumerate(NOTES, start=1):
        lines.append(f"## Slide {i:02d} - {title} ({timing})")
        lines.append(note)
        lines.append("")
    OUTPUT_NOTES.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh-screens", action="store_true", help="Recapture screenshots from the running local app.")
    args = parser.parse_args()

    ensure_dirs()
    collect_screenshots(refresh=args.refresh_screens)
    missing = [name for name in SCREENSHOT_URLS if not (SCREEN_DIR / name).exists()]
    if missing:
        raise SystemExit(f"Missing screenshots: {missing}")
    crops = make_crops()
    stats = fetch_live_stats()
    prs = build_deck(crops, stats)
    prs.save(OUTPUT_PPTX)
    write_notes(stats)
    print(json.dumps({
        "pptx": str(OUTPUT_PPTX),
        "notes": str(OUTPUT_NOTES),
        "assets": str(ASSET_DIR),
        "stats": stats.__dict__,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
