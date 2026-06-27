from __future__ import annotations

import datetime as dt
from collections import Counter
from io import BytesIO
from math import isfinite
from typing import Any, Iterable

from openpyxl import load_workbook

from .domain import AnnualMetricRow, CompanyMapping, FundamentalQualityIssue, FundamentalSnapshot, FundamentalWorkbook, PeriodMetricRow


CORE_SUMMARY_SHEET = "Factor_Summary_10Y"
LATEST_SUMMARY_SHEET = "Factor_Summary_Latest"
MARKET_MAP_SHEET = "Market_Map"
SUMMARY_SHEET = "Summary"

IGNORED_WIDE_SHEETS = {
    SUMMARY_SHEET,
    LATEST_SUMMARY_SHEET,
    CORE_SUMMARY_SHEET,
    "Coverage_10Y",
    "Data_Gap_Audit",
    "Long_Data",
    "Annual_Source",
    "prices",
    "nbre titres",
    MARKET_MAP_SHEET,
}

CORE_ID_COLUMNS = {"Company", "Statement_Year", "Has_Core_Fundamentals"}
LATEST_ID_COLUMNS = {"Company", "Latest Statement Year"}
LATEST_MARKET_ONLY_COLUMNS = {
    "Current_Price",
    "EnterpriseValue",
    "MarketCap_Calc",
    "Shares_Outstanding",
}


def _normalize_company(value: Any) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _normalize_symbol(value: Any) -> str | None:
    symbol = str(value or "").strip().upper()
    return symbol or None


def _safe_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if isfinite(out) else None


def _safe_int(value: Any) -> int | None:
    f = _safe_float(value)
    if f is None:
        return None
    return int(f)


def _rows_as_dicts(wb: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    out: list[dict[str, Any]] = []
    for row in rows:
        if not row or not any(v is not None for v in row):
            continue
        out.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return out


def _parse_summary(wb: Any) -> dict[str, Any]:
    rows = _rows_as_dicts(wb, SUMMARY_SHEET)
    return dict(rows[0]) if rows else {}


def _parse_mappings(wb: Any) -> list[CompanyMapping]:
    rows = _rows_as_dicts(wb, MARKET_MAP_SHEET)
    symbols = [
        _normalize_symbol(row.get("Ticker"))
        for row in rows
        if _normalize_symbol(row.get("Ticker")) is not None
    ]
    duplicate_symbols = {symbol for symbol, count in Counter(symbols).items() if count > 1}

    mappings: list[CompanyMapping] = []
    for row in rows:
        company_name = str(row.get("Company in Long_Data") or "").strip()
        if not company_name:
            continue
        symbol = _normalize_symbol(row.get("Ticker"))
        mappings.append(
            CompanyMapping(
                company_name=company_name,
                mapped_company_name=str(row.get("Mapped market company") or "").strip() or None,
                symbol=symbol,
                shares_outstanding=_safe_float(row.get("Shares outstanding")),
                match_type=str(row.get("Match type") or "").strip() or None,
                score_note=str(row.get("Score / Note") or "").strip() or None,
                source=str(row.get("Source") or "").strip() or None,
                canonical_company_name=_normalize_company(company_name),
                is_duplicate_symbol=bool(symbol and symbol in duplicate_symbols),
            )
        )
    return mappings


def _company_to_symbol(mappings: Iterable[CompanyMapping]) -> dict[str, str]:
    out: dict[str, str] = {}
    for mapping in mappings:
        if mapping.symbol:
            out[_normalize_company(mapping.company_name)] = mapping.symbol
    return out


def _company_to_shares(mappings: Iterable[CompanyMapping]) -> dict[str, float]:
    out: dict[str, float] = {}
    for mapping in mappings:
        if mapping.symbol and mapping.shares_outstanding and mapping.shares_outstanding > 0:
            out[mapping.symbol] = mapping.shares_outstanding
    return out


def _nonnull_metric_count(row: dict[str, Any], skip: set[str]) -> int:
    return sum(_safe_float(value) is not None for key, value in row.items() if key not in skip)


def _latest_fundamental_metric_count(row: dict[str, Any]) -> int:
    return sum(
        _safe_float(value) is not None
        for key, value in row.items()
        if key not in LATEST_ID_COLUMNS and key not in LATEST_MARKET_ONLY_COLUMNS
    )


def _select_best_latest(rows: list[dict[str, Any]], company_symbols: dict[str, str]) -> list[dict[str, Any]]:
    best: dict[str, tuple[tuple[int, int, int], dict[str, Any]]] = {}
    for row in rows:
        symbol = company_symbols.get(_normalize_company(row.get("Company")))
        if not symbol:
            continue
        year = _safe_int(row.get("Latest Statement Year")) or -1
        count = _nonnull_metric_count(row, LATEST_ID_COLUMNS)
        fundamental_count = _latest_fundamental_metric_count(row)
        key = (1 if fundamental_count > 0 else 0, year, count)
        if symbol not in best or key > best[symbol][0]:
            best[symbol] = (key, row)
    return [row for _, row in best.values()]


def _parse_latest_snapshots(
    wb: Any,
    mappings: list[CompanyMapping],
    annual_metrics: list[AnnualMetricRow],
) -> list[FundamentalSnapshot]:
    rows = _select_best_latest(_rows_as_dicts(wb, LATEST_SUMMARY_SHEET), _company_to_symbol(mappings))
    company_symbols = _company_to_symbol(mappings)
    shares_by_symbol = _company_to_shares(mappings)
    annual_count_by_symbol: dict[str, int] = {}
    for row in annual_metrics:
        if row.metric_name == "Has_Core_Fundamentals" and row.metric_value and row.metric_value > 0:
            annual_count_by_symbol[row.symbol] = annual_count_by_symbol.get(row.symbol, 0) + 1

    snapshots: list[FundamentalSnapshot] = []
    for row in rows:
        company_name = str(row.get("Company") or "").strip()
        symbol = company_symbols.get(_normalize_company(company_name))
        if not symbol:
            continue
        metrics = {
            key: _safe_float(value)
            for key, value in row.items()
            if key not in LATEST_ID_COLUMNS and key
        }
        if shares_by_symbol.get(symbol):
            metrics.setdefault("Shares_Outstanding", shares_by_symbol[symbol])
            market_cap = metrics.get("MarketCap_Calc")
            if market_cap and market_cap > 0:
                metrics.setdefault("Current_Price", market_cap / shares_by_symbol[symbol])
        coverage = {
            "latest_statement_year": _safe_int(row.get("Latest Statement Year")),
            "metric_count": sum(value is not None for value in metrics.values()),
            "core_year_count": annual_count_by_symbol.get(symbol, 0),
            "has_market_cap": metrics.get("MarketCap_Calc") is not None,
            "has_current_price": metrics.get("Current_Price") is not None,
        }
        snapshots.append(
            FundamentalSnapshot(
                symbol=symbol,
                company_name=company_name,
                latest_statement_year=_safe_int(row.get("Latest Statement Year")),
                metrics=metrics,
                coverage=coverage,
                source={
                    "workbook_sheet": LATEST_SUMMARY_SHEET,
                    "deduped_by": "latest_statement_year_then_metric_count",
                    "lineage": {
                        key: {"sheet": LATEST_SUMMARY_SHEET, "field": key}
                        for key in metrics
                    },
                },
            )
        )
    return sorted(snapshots, key=lambda item: item.symbol)


def _parse_core_summary_annual(wb: Any, mappings: list[CompanyMapping]) -> list[AnnualMetricRow]:
    rows = _rows_as_dicts(wb, CORE_SUMMARY_SHEET)
    company_symbols = _company_to_symbol(mappings)
    best: dict[tuple[str, int], tuple[int, dict[str, Any]]] = {}
    for row in rows:
        symbol = company_symbols.get(_normalize_company(row.get("Company")))
        year = _safe_int(row.get("Statement_Year"))
        if not symbol or year is None:
            continue
        count = _nonnull_metric_count(row, {"Company", "Statement_Year"})
        if count == 0:
            continue
        key = (symbol, year)
        if key not in best or count > best[key][0]:
            best[key] = (count, row)

    out: list[AnnualMetricRow] = []
    for (symbol, year), (_, row) in sorted(best.items()):
        company_name = str(row.get("Company") or "").strip()
        for metric_name, value in row.items():
            if metric_name in {"Company", "Statement_Year"} or not metric_name:
                continue
            out.append(
                AnnualMetricRow(
                    symbol=symbol,
                    company_name=company_name,
                    statement_year=year,
                    metric_name=metric_name,
                    metric_value=_safe_float(value),
                    raw_metric_name=metric_name,
                    source_sheet=CORE_SUMMARY_SHEET,
                    source_field=metric_name,
                )
            )
    return out


def _sheet_looks_wide_metric(headers: list[Any]) -> bool:
    if not headers or str(headers[0]).strip() != "Company":
        return False
    year_count = 0
    for value in headers[1:]:
        year = _safe_int(value)
        if year is not None and 1900 <= year <= 2100:
            year_count += 1
    return year_count >= 2


def _parse_wide_metric_sheets(wb: Any, mappings: list[CompanyMapping]) -> list[AnnualMetricRow]:
    company_symbols = _company_to_symbol(mappings)
    out: list[AnnualMetricRow] = []
    seen: set[tuple[str, int, str]] = set()
    for sheet_name in wb.sheetnames:
        if sheet_name in IGNORED_WIDE_SHEETS:
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        if not _sheet_looks_wide_metric(headers):
            continue
        year_columns = [(index, _safe_int(value)) for index, value in enumerate(headers) if _safe_int(value) is not None]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            company_name = str(row[0]).strip()
            symbol = company_symbols.get(_normalize_company(company_name))
            if not symbol:
                continue
            for index, year in year_columns:
                if year is None or index >= len(row):
                    continue
                key = (symbol, year, sheet_name)
                if key in seen:
                    continue
                value = _safe_float(row[index])
                if value is None:
                    continue
                seen.add(key)
                out.append(
                    AnnualMetricRow(
                        symbol=symbol,
                        company_name=company_name,
                        statement_year=year,
                        metric_name=sheet_name,
                        metric_value=value,
                        raw_metric_name=sheet_name,
                        source_sheet=sheet_name,
                        source_field=str(year),
                    )
                )
    return out


def _merge_annual_metrics(rows: Iterable[AnnualMetricRow]) -> list[AnnualMetricRow]:
    best: dict[tuple[str, int, str], AnnualMetricRow] = {}
    for row in rows:
        key = (row.symbol, row.statement_year, row.metric_name)
        if key not in best or (best[key].metric_value is None and row.metric_value is not None):
            best[key] = row
    return sorted(best.values(), key=lambda item: (item.symbol, item.statement_year, item.metric_name))


def _annual_rows_as_period_metrics(rows: Iterable[AnnualMetricRow]) -> list[PeriodMetricRow]:
    return [
        PeriodMetricRow(
            symbol=row.symbol,
            company_name=row.company_name,
            fiscal_year=row.statement_year,
            period_type="annual",
            period_label="FY",
            period_end_date=f"{row.statement_year:04d}-12-31",
            metric_name=row.metric_name,
            metric_value=row.metric_value,
            raw_metric_name=row.raw_metric_name,
            source_url=None,
            document_title=row.source_sheet,
            is_proxy=row.is_proxy,
        )
        for row in rows
        if row.metric_value is not None
    ]


def _build_quality_issues(
    mappings: list[CompanyMapping],
    annual_metrics: list[AnnualMetricRow],
    latest_snapshots: list[FundamentalSnapshot],
) -> list[FundamentalQualityIssue]:
    issues: list[FundamentalQualityIssue] = []
    now_year = dt.date.today().year
    mapped_symbols = {mapping.symbol for mapping in mappings if mapping.symbol}
    duplicate_symbols = sorted({mapping.symbol for mapping in mappings if mapping.symbol and mapping.is_duplicate_symbol})
    for symbol in duplicate_symbols:
        issues.append(
            FundamentalQualityIssue(
                severity="warning",
                code="duplicate_symbol_mapping",
                message=f"Multiple workbook company rows map to {symbol}; latest snapshot is deduped by year and metric coverage.",
                symbol=symbol,
            )
        )
    if not mapped_symbols:
        issues.append(
            FundamentalQualityIssue(
                severity="error",
                code="no_mapped_symbols",
                message="The workbook did not provide any usable ticker mapping.",
            )
        )
    annual_by_symbol: dict[str, list[AnnualMetricRow]] = {}
    for row in annual_metrics:
        annual_by_symbol.setdefault(row.symbol, []).append(row)
        if row.statement_year > now_year + 1:
            issues.append(
                FundamentalQualityIssue(
                    severity="warning",
                    code="future_statement_year",
                    message="Statement year is beyond the expected current reporting window.",
                    symbol=row.symbol,
                    metric_name=row.metric_name,
                    statement_year=row.statement_year,
                )
            )
    for snapshot in latest_snapshots:
        core_year_count = int(snapshot.coverage.get("core_year_count") or 0)
        if core_year_count < 3:
            issues.append(
                FundamentalQualityIssue(
                    severity="warning",
                    code="short_fundamental_history",
                    message="Less than three years of core fundamentals are available; trend models will have lower confidence.",
                    symbol=snapshot.symbol,
                    context={"core_year_count": core_year_count},
                )
            )
        if not snapshot.coverage.get("has_current_price"):
            issues.append(
                FundamentalQualityIssue(
                    severity="warning",
                    code="missing_current_price_proxy",
                    message="Current price could not be derived from market cap and shares outstanding.",
                    symbol=snapshot.symbol,
                )
            )
        if snapshot.symbol not in annual_by_symbol:
            issues.append(
                FundamentalQualityIssue(
                    severity="warning",
                    code="missing_annual_history",
                    message="Latest snapshot exists without normalized annual metric history.",
                    symbol=snapshot.symbol,
                )
            )
    return issues


def parse_fundamental_workbook(payload: bytes) -> FundamentalWorkbook:
    """Parse the structured market-formula workbook into normalized rows.

    The workbook stores formula results as cached Excel values. This parser uses
    ``data_only=True`` and keeps missing formulas as ``None`` so downstream code
    can mark model confidence honestly.
    """

    wb = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    mappings = _parse_mappings(wb)
    core_rows = _parse_core_summary_annual(wb, mappings)
    wide_rows = _parse_wide_metric_sheets(wb, mappings)
    annual_metrics = _merge_annual_metrics([*core_rows, *wide_rows])
    latest_snapshots = _parse_latest_snapshots(wb, mappings, annual_metrics)
    quality_issues = _build_quality_issues(mappings, annual_metrics, latest_snapshots)

    mapped_symbols = {mapping.symbol for mapping in mappings if mapping.symbol}
    summary = {
        **_parse_summary(wb),
        "mapped_company_count": len(mappings),
        "mapped_symbol_count": len(mapped_symbols),
        "annual_metric_count": len(annual_metrics),
        "latest_snapshot_count": len(latest_snapshots),
        "duplicate_symbols": sorted(
            {
                mapping.symbol
                for mapping in mappings
                if mapping.symbol and mapping.is_duplicate_symbol
            }
        ),
        "quality_issue_count": len(quality_issues),
        "quality_issue_counts": dict(Counter(issue.severity for issue in quality_issues)),
    }
    return FundamentalWorkbook(
        mappings=mappings,
        annual_metrics=annual_metrics,
        latest_snapshots=latest_snapshots,
        summary=summary,
        quality_issues=quality_issues,
        period_metrics=_annual_rows_as_period_metrics(annual_metrics),
    )
